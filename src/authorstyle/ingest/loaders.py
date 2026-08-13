from __future__ import annotations

from abc import ABC, abstractmethod
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument

from authorstyle.ingest.text_utils import normalize_unicode, parse_front_matter


class LoadedDocument:
    def __init__(
        self,
        source_path: Path,
        raw_text: str,
        metadata: dict,
        headings: list[str] | None = None,
    ) -> None:
        self.source_path = source_path
        self.raw_text = raw_text
        self.metadata = metadata
        self.headings = headings or []


class BaseDocumentLoader(ABC):
    extensions: tuple[str, ...] = ()

    @abstractmethod
    def load(self, path: Path) -> LoadedDocument:
        ...


class TextLoader(BaseDocumentLoader):
    extensions = (".txt",)

    def load(self, path: Path) -> LoadedDocument:
        raw = path.read_text(encoding="utf-8", errors="replace")
        return LoadedDocument(path, normalize_unicode(raw), {})


class MarkdownLoader(BaseDocumentLoader):
    extensions = (".md",)

    def load(self, path: Path) -> LoadedDocument:
        raw = path.read_text(encoding="utf-8", errors="replace")
        meta, body = parse_front_matter(raw)
        headings = [
            line.lstrip("#").strip()
            for line in body.splitlines()
            if line.startswith("#")
        ]
        return LoadedDocument(path, normalize_unicode(body), meta, headings)


class HtmlLoader(BaseDocumentLoader):
    extensions = (".html", ".htm")

    def load(self, path: Path) -> LoadedDocument:
        raw = path.read_text(encoding="utf-8", errors="replace")
        soup = BeautifulSoup(raw, "html.parser")
        title_tag = soup.find("title")
        meta = {}
        if title_tag and title_tag.string:
            meta["title"] = title_tag.string.strip()
        text = soup.get_text("\n")
        return LoadedDocument(path, normalize_unicode(text), meta)


class DocxLoader(BaseDocumentLoader):
    extensions = (".docx",)

    def load(self, path: Path) -> LoadedDocument:
        doc = DocxDocument(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        meta = {}
        if doc.core_properties.title:
            meta["title"] = doc.core_properties.title
        return LoadedDocument(path, normalize_unicode("\n\n".join(paragraphs)), meta)


LOADERS: list[BaseDocumentLoader] = [
    MarkdownLoader(),
    TextLoader(),
    HtmlLoader(),
    DocxLoader(),
]


def loader_for(path: Path) -> BaseDocumentLoader | None:
    suffix = path.suffix.lower()
    for loader in LOADERS:
        if suffix in loader.extensions:
            return loader
    return None
